#!/usr/bin/env python3
"""Build RoboForeman 3-minute pitch deck v2 — minimal, Keynote-friendly PPTX.

Keynote compatibility notes:
- Mac-native fonts only (Helvetica Neue / Helvetica) — Calibri is Office-only
- Plain rectangles instead of rounded-rect adjustments (import more reliably)
- Standard 16:9 widescreen, embedded PNG assets, speaker notes preserved
- Open in Keynote → File → Save to get a native .key on Mac
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "RoboForeman-Pitch-3min-v2.pptx"
# Alias filename so it's obvious this PPTX is meant for Keynote on Mac
OUT_KEYNOTE = ROOT / "RoboForeman-Pitch-3min-v2.keynote.pptx"

# ── Minimal palette (no purple) ─────────────────────────────────────────────
INK = RGBColor(0x14, 0x14, 0x14)
MUTED = RGBColor(0x5C, 0x5C, 0x5C)
FAINT = RGBColor(0x9A, 0x9A, 0x9A)
RULE = RGBColor(0xE8, 0xE8, 0xE8)
BG = RGBColor(0xFF, 0xFF, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xC2, 0x6A, 0x00)  # construction amber — not purple
AMBER_DEEP = RGBColor(0x9A, 0x54, 0x00)
CHAR = RGBColor(0x16, 0x16, 0x16)
SOFT = RGBColor(0xF2, 0xF2, 0xF2)  # cool gray panel, not cream
OK = RGBColor(0x2F, 0x6B, 0x3A)
BAD = RGBColor(0xA8, 0x3A, 0x2C)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Mac-native fonts — Keynote maps these without substitution
FONT_DISPLAY = "Helvetica Neue"
FONT_BODY = "Helvetica"


def set_run_font(run, size, bold=False, color=INK, font_name=FONT_BODY, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = etree.SubElement(rPr, qn(f"a:{tag}"))
        el.set("typeface", font_name)


def add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def write_para(tf, text, size, bold=False, color=INK, font_name=FONT_BODY,
               align=PP_ALIGN.LEFT, space_before=0, space_after=0, italic=False):
    p = tf.paragraphs[0] if not tf.paragraphs[0].text else tf.add_paragraph()
    if not tf.paragraphs[0].text and len(tf.paragraphs) == 1:
        p = tf.paragraphs[0]
    else:
        # first paragraph may already have content; ensure we use empty first once
        if tf.paragraphs[0].text == "" and len(tf.paragraphs) == 1:
            p = tf.paragraphs[0]
    p.clear()
    run = p.add_run()
    run.text = text
    set_run_font(run, size, bold=bold, color=color, font_name=font_name, italic=italic)
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    return p


def set_text(shape, lines, *, default_size=18, default_bold=False, default_color=INK,
             default_font=FONT_BODY, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    """lines: list of str or dicts with text/size/bold/color/font/italic/space_before/space_after."""
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    try:
        tf.auto_size = None
    except Exception:
        pass
    shape.text_frame.paragraphs[0].alignment = align

    # vertical anchor
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    if bodyPr is not None:
        anchor_map = {
            MSO_ANCHOR.TOP: "t",
            MSO_ANCHOR.MIDDLE: "ctr",
            MSO_ANCHOR.BOTTOM: "b",
        }
        bodyPr.set("anchor", anchor_map.get(valign, "t"))

    first = True
    for line in lines:
        if isinstance(line, str):
            line = {"text": line}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = line.get("align", align)
        p.space_before = Pt(line.get("space_before", 0))
        p.space_after = Pt(line.get("space_after", 4))
        run = p.add_run()
        run.text = line["text"]
        set_run_font(
            run,
            line.get("size", default_size),
            bold=line.get("bold", default_bold),
            color=line.get("color", default_color),
            font_name=line.get("font", default_font),
            italic=line.get("italic", False),
        )


def add_rect(slide, left, top, width, height, fill=None, line=None, line_width_pt=1):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_width_pt)
    return shape


def add_round_rect(slide, left, top, width, height, fill=None, line=None, line_width_pt=1):
    """Use a plain rectangle — Keynote imports these more reliably than adjusted rounded rects."""
    return add_rect(slide, left, top, width, height, fill=fill, line=line, line_width_pt=line_width_pt)


def fill_slide_bg(slide, color=BG):
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=color)


def accent_bar(slide, left=Inches(0.7), top=Inches(0.55), width=Inches(0.55), height=Inches(0.08)):
    add_rect(slide, left, top, width, height, fill=AMBER)


def footer(slide, page, total=7):
    box = add_textbox(slide, Inches(0.7), Inches(7.05), Inches(8), Inches(0.3))
    set_text(
        box,
        [{"text": "roboforeman  ·  SmartSense  ·  Cursor Hackathon", "size": 11, "color": FAINT, "font": FONT_BODY}],
    )
    num = add_textbox(slide, Inches(11.8), Inches(7.05), Inches(1.0), Inches(0.3))
    set_text(
        num,
        [{"text": f"{page:02d} / {total:02d}", "size": 11, "color": FAINT, "align": PP_ALIGN.RIGHT}],
        align=PP_ALIGN.RIGHT,
    )


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


# ── Slides ──────────────────────────────────────────────────────────────────

def slide_1_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_slide_bg(slide, CHAR)

    # thin amber top rule
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.08), fill=AMBER)

    # mark
    if (ASSETS / "hardhat-mark.png").exists():
        slide.shapes.add_picture(
            str(ASSETS / "hardhat-mark.png"),
            Inches(0.85), Inches(1.55),
            Inches(0.72), Inches(0.72),
        )

    eyebrow = add_textbox(slide, Inches(0.85), Inches(2.45), Inches(11), Inches(0.35))
    set_text(
        eyebrow,
        [{"text": "TEAM SMARTSENSE  ·  AI FOR EVERYDAY EXPERIENCES", "size": 12, "bold": True, "color": AMBER, "font": FONT_BODY}],
    )

    title = add_textbox(slide, Inches(0.85), Inches(2.9), Inches(11.5), Inches(1.1))
    set_text(
        title,
        [{"text": "roboforeman", "size": 60, "bold": True, "color": WHITE, "font": FONT_DISPLAY}],
    )

    line = add_textbox(slide, Inches(0.85), Inches(4.15), Inches(11.2), Inches(1.2))
    set_text(
        line,
        [
            {
                "text": "Your AI builds anything — it just doesn't know your project.",
                "size": 26,
                "color": RGBColor(0xD0, 0xD0, 0xC8),
                "font": FONT_DISPLAY,
                "italic": True,
            }
        ],
    )

    sub = add_textbox(slide, Inches(0.85), Inches(5.55), Inches(11), Inches(0.4))
    set_text(
        sub,
        [{"text": "One scan. Agent-ready. Ready to contribute.", "size": 16, "color": FAINT, "font": FONT_BODY}],
    )

    notes(
        slide,
        """AI can build almost anything now. But it doesn't know your project.
Roboforeman fixes that — in two minutes. Let me show you.""",
    )


def slide_2_pain(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_slide_bg(slide)
    accent_bar(slide)

    label = add_textbox(slide, Inches(0.7), Inches(0.75), Inches(4), Inches(0.3))
    set_text(label, [{"text": "THE PROBLEM", "size": 12, "bold": True, "color": AMBER}])

    title = add_textbox(slide, Inches(0.7), Inches(1.15), Inches(11.8), Inches(1.0))
    set_text(
        title,
        [
            {
                "text": "Meet Priya. Shipping a new feature.",
                "size": 36,
                "bold": True,
                "color": INK,
                "font": FONT_DISPLAY,
            }
        ],
    )

    # three failure chips — minimal, not card-heavy
    fails = [
        ("Wrong style", "Ignores how her project actually writes code"),
        ("Fake commands", "Calls scripts that don't exist in the repo"),
        ("Breaks working code", "Touches files it shouldn't — now she fixes the AI"),
    ]
    x0 = Inches(0.7)
    gap = Inches(0.35)
    w = Inches(3.85)
    for i, (head, body) in enumerate(fails):
        left = x0 + i * (w + gap)
        add_rect(slide, left, Inches(2.7), Inches(0.08), Inches(2.2), fill=BAD)
        h = add_textbox(slide, left + Inches(0.3), Inches(2.75), w - Inches(0.3), Inches(0.55))
        set_text(h, [{"text": head, "size": 20, "bold": True, "color": INK, "font": FONT_DISPLAY}])
        b = add_textbox(slide, left + Inches(0.3), Inches(3.4), w - Inches(0.35), Inches(1.3))
        set_text(b, [{"text": body, "size": 15, "color": MUTED}])

    bottom = add_textbox(slide, Inches(0.7), Inches(5.6), Inches(11.8), Inches(0.7))
    set_text(
        bottom,
        [
            {
                "text": "She's correcting the assistant instead of building. Every developer here has lived this.",
                "size": 18,
                "italic": True,
                "color": INK,
                "font": FONT_DISPLAY,
            }
        ],
    )
    footer(slide, 2)
    notes(
        slide,
        """This is Priya. She's got the smartest assistant in the world in her editor.
She says "add the payment screen" — and it writes code that ignores her project's style,
calls commands that don't exist, and breaks something that worked.
Now she's fixing the AI instead of building. Every developer here has lived this.""",
    )


def slide_3_reframe(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_slide_bg(slide)
    accent_bar(slide)

    label = add_textbox(slide, Inches(0.7), Inches(0.75), Inches(4), Inches(0.3))
    set_text(label, [{"text": "THE REFRAME", "size": 12, "bold": True, "color": AMBER}])

    title = add_textbox(slide, Inches(0.7), Inches(1.15), Inches(11.8), Inches(1.1))
    set_text(
        title,
        [
            {
                "text": "The AI isn't dumb — it's unbriefed.",
                "size": 36,
                "bold": True,
                "color": INK,
                "font": FONT_DISPLAY,
            }
        ],
    )

    # metaphor mapping — clean two columns
    left_items = [
        ("No site map", "No AGENTS.md"),
        ("No rulebook", "No .cursor/rules"),
        ("No safety tape", "No hooks / .cursorignore"),
        ("No crew habits", "No skills from your code"),
    ]

    for i, (story, reality) in enumerate(left_items):
        top = Inches(2.55) + i * Inches(0.7)
        a = add_textbox(slide, Inches(0.7), top, Inches(5.2), Inches(0.55))
        set_text(
            a,
            [
                {"text": story, "size": 20, "bold": True, "color": INK, "font": FONT_DISPLAY, "space_after": 0},
            ],
        )
        b = add_textbox(slide, Inches(6.2), top, Inches(6.2), Inches(0.55))
        set_text(b, [{"text": reality, "size": 18, "color": MUTED}])

    # punchline strip
    add_rect(slide, Inches(0.7), Inches(5.55), Inches(11.9), Inches(0.9), fill=CHAR)
    punch = add_textbox(slide, Inches(1.0), Inches(5.7), Inches(11.3), Inches(0.6))
    set_text(
        punch,
        [
            {
                "text": "Roboforeman briefs it — so the builder stops guessing.",
                "size": 22,
                "bold": True,
                "color": WHITE,
                "font": FONT_DISPLAY,
            }
        ],
        valign=MSO_ANCHOR.MIDDLE,
    )
    footer(slide, 3)
    notes(
        slide,
        """But the AI isn't dumb — it's unbriefed. Drop a world-class builder on a site and tell them
nothing, of course they guess wrong. That's Roboforeman's job: it walks your codebase, sees
what's missing, and hands the AI the site map, the rulebook, the safety tape, and how your team
actually does things. Watch it run.""",
    )


def slide_4_demo(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_slide_bg(slide, CHAR)
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.08), fill=AMBER)

    label = add_textbox(slide, Inches(0.7), Inches(0.35), Inches(6), Inches(0.3))
    set_text(label, [{"text": "DEMO  ·  6× FAST-FORWARD", "size": 12, "bold": True, "color": AMBER}])

    title = add_textbox(slide, Inches(0.7), Inches(0.7), Inches(12), Inches(0.6))
    set_text(
        title,
        [
            {
                "text": "/scan  →  /build  →  /contribute",
                "size": 32,
                "bold": True,
                "color": WHITE,
                "font": FONT_DISPLAY,
            }
        ],
    )

    # large video frame
    frame_left, frame_top = Inches(0.7), Inches(1.5)
    frame_w, frame_h = Inches(8.6), Inches(5.2)
    add_round_rect(slide, frame_left, frame_top, frame_w, frame_h, fill=RGBColor(0x2A, 0x2A, 0x28), line=RGBColor(0x3A, 0x3A, 0x38))

    if (ASSETS / "play-badge.png").exists():
        slide.shapes.add_picture(
            str(ASSETS / "play-badge.png"),
            Inches(4.35), Inches(3.4),
            Inches(1.1), Inches(1.1),
        )

    hint = add_textbox(slide, Inches(1.1), Inches(4.7), Inches(7.8), Inches(0.8))
    set_text(
        hint,
        [
            {
                "text": "Drop your demo video here\nReplace this frame before the pitch",
                "size": 14,
                "color": FAINT,
                "align": PP_ALIGN.CENTER,
            }
        ],
        align=PP_ALIGN.CENTER,
    )

    # right rail — what judges will see
    steps = [
        ("01  /scan", "Stack detect + honest checklist"),
        ("02  /build", "Rules, map, skills — backed up first"),
        ("03  /contribute", "Issue → fix → draft PR, no pollution"),
    ]
    y = Inches(1.55)
    for head, body in steps:
        add_rect(slide, Inches(9.6), y, Inches(0.08), Inches(1.25), fill=AMBER)
        h = add_textbox(slide, Inches(9.9), y, Inches(2.9), Inches(0.4))
        set_text(h, [{"text": head, "size": 14, "bold": True, "color": AMBER}])
        b = add_textbox(slide, Inches(9.9), y + Inches(0.4), Inches(2.9), Inches(0.7))
        set_text(b, [{"text": body, "size": 14, "color": RGBColor(0xC8, 0xC8, 0xC0)}])
        y += Inches(1.55)

    notes(
        slide,
        """One command — it inspects the repo with an honest checklist. No vibes.
Second command — I tick what I want, hit Lock it in, and the crew works: rules, a project map,
and skills learned from the repo's own code. Everything's backed up first — nothing overwritten.
Then /contribute — pick a good-first-issue, solve it to definition-of-done, ship a draft PR
with zero foreman files in the diff. Now the same request as before… and this time, the AI nails it.""",
    )


def slide_5_diff(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_slide_bg(slide)
    accent_bar(slide)

    label = add_textbox(slide, Inches(0.7), Inches(0.75), Inches(6), Inches(0.3))
    set_text(label, [{"text": "WHY NOT JUST PROMPT IT?", "size": 12, "bold": True, "color": AMBER}])

    title = add_textbox(slide, Inches(0.7), Inches(1.15), Inches(11.8), Inches(1.2))
    set_text(
        title,
        [
            {
                "text": '"Can\'t I just make these files myself?"',
                "size": 32,
                "bold": True,
                "color": INK,
                "font": FONT_DISPLAY,
                "italic": True,
            }
        ],
    )

    # left: the trap
    add_rect(slide, Inches(0.7), Inches(2.7), Inches(5.7), Inches(3.2), fill=SOFT)
    left_t = add_textbox(slide, Inches(1.05), Inches(3.0), Inches(5.0), Inches(2.6))
    set_text(
        left_t,
        [
            {"text": "Sure — once.", "size": 22, "bold": True, "color": INK, "font": FONT_DISPLAY, "space_after": 12},
            {"text": "For one repo.", "size": 18, "color": MUTED, "space_after": 8},
            {"text": "If you remember every rule,", "size": 18, "color": MUTED, "space_after": 8},
            {"text": "every time — and never commit a secret.", "size": 18, "color": MUTED},
        ],
    )

    # right: the answer
    add_rect(slide, Inches(6.85), Inches(2.7), Inches(5.7), Inches(3.2), fill=CHAR)
    right_t = add_textbox(slide, Inches(7.2), Inches(3.0), Inches(5.0), Inches(2.6))
    set_text(
        right_t,
        [
            {"text": "Roboforeman", "size": 22, "bold": True, "color": AMBER, "font": FONT_DISPLAY, "space_after": 14},
            {"text": "instantly  ·  identically", "size": 18, "color": WHITE, "space_after": 8},
            {"text": "safely  ·  on every repo", "size": 18, "color": WHITE, "space_after": 8},
            {"text": "including open-source you don't own", "size": 16, "color": FAINT},
        ],
    )

    footer(slide, 5)
    notes(
        slide,
        """Now engineers are thinking — I could just make those files myself. You could. Once. For one
repo. If you remember every rule every time and never commit a secret. Roboforeman does it
instantly, the same way, on every project — and never forgets the safety tape. On open-source,
it won't even pollute your PR. That's automation — not a prompt you retype.""",
    )


def slide_6_hood(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_slide_bg(slide)
    accent_bar(slide)

    label = add_textbox(slide, Inches(0.7), Inches(0.75), Inches(6), Inches(0.3))
    set_text(label, [{"text": "UNDER THE HOOD", "size": 12, "bold": True, "color": AMBER}])

    title = add_textbox(slide, Inches(0.7), Inches(1.15), Inches(11.8), Inches(0.9))
    set_text(
        title,
        [
            {
                "text": "Fun on top. Real engineering underneath.",
                "size": 32,
                "bold": True,
                "color": INK,
                "font": FONT_DISPLAY,
            }
        ],
    )

    # two pillars
    pillars = [
        (
            "AI for judgment",
            [
                "Reads your code — and your CONTRIBUTING.md",
                "Mines skills from real exemplar files",
                "Exa → live stack docs, not stale guesses",
            ],
        ),
        (
            "Deterministic for trust",
            [
                "Tested detection & merge scripts",
                "Never-clobber + one-command undo",
                "No-pollution guards on OSS PRs",
            ],
        ),
    ]
    for i, (head, bullets) in enumerate(pillars):
        left = Inches(0.7) + i * Inches(6.2)
        add_rect(slide, left, Inches(2.35), Inches(0.08), Inches(2.8), fill=AMBER)
        h = add_textbox(slide, left + Inches(0.3), Inches(2.35), Inches(5.5), Inches(0.5))
        set_text(h, [{"text": head, "size": 20, "bold": True, "color": INK, "font": FONT_DISPLAY}])
        for j, b in enumerate(bullets):
            t = add_textbox(slide, left + Inches(0.3), Inches(3.05) + j * Inches(0.55), Inches(5.5), Inches(0.5))
            set_text(t, [{"text": b, "size": 16, "color": MUTED}])

    # sponsors — quiet strip
    add_rect(slide, Inches(0.7), Inches(5.6), Inches(11.9), Inches(0.85), fill=SOFT)
    sp = add_textbox(slide, Inches(1.0), Inches(5.75), Inches(11.3), Inches(0.55))
    set_text(
        sp,
        [
            {
                "text": "Exa  ·  ElevenLabs (the foreman speaks)  ·  Render",
                "size": 16,
                "bold": True,
                "color": INK,
            }
        ],
        valign=MSO_ANCHOR.MIDDLE,
    )

    footer(slide, 6)
    notes(
        slide,
        """Under the cartoon it's a real, standard Cursor plugin. The decisions are AI — it reads your
code and, with Exa, your stack's live docs. The plumbing is deterministic — tested detection,
never overwrites your files, one-command undo, and triple protection so foreman files never
land in an OSS PR. We even gave the foreman a voice with ElevenLabs.
Judgment where you want AI; trust where you want safety.""",
    )


def slide_7_close(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_slide_bg(slide, CHAR)
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.08), fill=AMBER)

    if (ASSETS / "hardhat-mark.png").exists():
        slide.shapes.add_picture(
            str(ASSETS / "hardhat-mark.png"),
            Inches(6.25), Inches(1.35),
            Inches(0.85), Inches(0.85),
        )

    title = add_textbox(slide, Inches(0.8), Inches(2.5), Inches(11.7), Inches(1.2))
    set_text(
        title,
        [
            {
                "text": "Can we fix your repo? Yes we can.",
                "size": 36,
                "bold": True,
                "color": WHITE,
                "font": FONT_DISPLAY,
                "align": PP_ALIGN.CENTER,
            }
        ],
        align=PP_ALIGN.CENTER,
    )

    sub = add_textbox(slide, Inches(1.5), Inches(3.85), Inches(10.3), Inches(0.9))
    set_text(
        sub,
        [
            {
                "text": "Brief your AI. Build your features.\nContribute to open source — without the setup tax.",
                "size": 18,
                "color": RGBColor(0xC8, 0xC8, 0xC0),
                "align": PP_ALIGN.CENTER,
            }
        ],
        align=PP_ALIGN.CENTER,
    )

    brand = add_textbox(slide, Inches(0.8), Inches(5.3), Inches(11.7), Inches(0.5))
    set_text(
        brand,
        [
            {
                "text": "roboforeman",
                "size": 20,
                "bold": True,
                "color": AMBER,
                "font": FONT_DISPLAY,
                "align": PP_ALIGN.CENTER,
            }
        ],
        align=PP_ALIGN.CENTER,
    )

    thanks = add_textbox(slide, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.4))
    set_text(
        thanks,
        [
            {
                "text": "SmartSense  ·  Built on Cursor  ·  Thank you",
                "size": 13,
                "color": FAINT,
                "align": PP_ALIGN.CENTER,
            }
        ],
        align=PP_ALIGN.CENTER,
    )

    notes(
        slide,
        """AI can build anything — it just needed to know your project. Now it does.
Point it at your repo or any open-source giant — scan, build, contribute.
Can we fix your repo? Yes we can. Thank you.""",
    )


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_1_title(prs)
    slide_2_pain(prs)
    slide_3_reframe(prs)
    slide_4_demo(prs)
    slide_5_diff(prs)
    slide_6_hood(prs)
    slide_7_close(prs)

    prs.save(OUT)
    # Same bytes under a clearer Mac/Keynote-facing name
    prs.save(OUT_KEYNOTE)
    print(f"Wrote {OUT}")
    print(f"Wrote {OUT_KEYNOTE}")


if __name__ == "__main__":
    main()
